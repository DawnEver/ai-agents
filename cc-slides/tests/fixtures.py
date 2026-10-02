"""Synthetic decks for tests — no real (branded / confidential) template is committed."""
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Emu, Pt

W, H = 12192000, 6858000  # 16:9
# Three-up pattern: two boxes on top, one centred below (reading order TL, TR, B).
BOX_TL = (58213, 905386, 4877334, 3427120)
BOX_B = (3292985, 3589776, 5239686, 3285460)
BOX_TR = (6353541, 900417, 5780246, 3437059)
CAPTION = (500000, 6400000, 4000000, 400000)


def make_image(path, size, color=(200, 60, 60)):
    Image.new("RGB", size, color).save(path)
    return Path(path)


def make_template(directory):
    """Slide 1: cover (title only). Slide 2: three-up pattern with caption and a static line."""
    directory = Path(directory)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    layout = prs.slide_layouts[5]  # Title Only

    cover = prs.slides.add_slide(layout)
    cover.shapes.title.text = "Cover"

    pattern = prs.slides.add_slide(layout)
    pattern.shapes.title.text = "Pattern title"
    placeholder = make_image(directory / "placeholder.png", (400, 300))
    # Insert out of reading order on purpose: TL, B, TR (as in the real reference deck).
    for name, box in (("Picture TL", BOX_TL), ("Picture B", BOX_B), ("Picture TR", BOX_TR)):
        pic = pattern.shapes.add_picture(str(placeholder), *map(Emu, box))
        pic.name = name
    caption = pattern.shapes.add_textbox(*map(Emu, CAPTION))
    caption.name = "Caption"
    caption.text_frame.text = "Old caption"
    caption.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    line = pattern.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(0), Emu(880000), Emu(W), Emu(880000))
    line.name = "Rule"

    path = directory / "template.pptx"
    prs.save(path)
    return path

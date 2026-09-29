import hashlib
import re
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

import build_slides  # noqa: E402
import slides_lib  # noqa: E402
from fixtures import BOX_B, BOX_TL, BOX_TR, make_image, make_template  # noqa: E402

SPEC = """
template = "template.pptx"

[patterns.three-up]
slide = 2

[[slides]]
keep = 1

[[slides]]
pattern = "three-up"
title = "Efficiency Map"
images = ["wide.png", "tall.png", "square.png"]
texts = ["New caption"]
notes = "Speaker note"

[[slides]]
pattern = "three-up"
title = "Loss Map"
images = [{ path = "wide.png", fit = "cover" }, "", "square.png"]
"""


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def inside(shape, box, tol=2):
    x, y, w, h = box
    return (shape.left >= x - tol and shape.top >= y - tol
            and shape.left + shape.width <= x + w + tol and shape.top + shape.height <= y + h + tol)


class BuildTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.template = make_template(self.dir)
        make_image(self.dir / "wide.png", (1600, 900))
        make_image(self.dir / "tall.png", (600, 1200))
        make_image(self.dir / "square.png", (800, 800))
        self.spec = self.dir / "deck.toml"
        self.spec.write_text(SPEC)

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, text=None):
        if text is not None:
            self.spec.write_text(text)
        out = self.dir / "out.pptx"
        build_slides.build(self.spec, out)
        return Presentation(out)

    def test_slide_order_keep_and_pattern_removed(self):
        before = sha(self.template)
        prs = self.build()
        titles = [s.shapes.title.text for s in prs.slides]
        self.assertEqual(titles, ["Cover", "Efficiency Map", "Loss Map"])
        self.assertEqual(sha(self.template), before, "template must never be modified")

    def test_images_follow_reading_order_and_keep_aspect(self):
        slide = self.build().slides[1]
        pics = {p.name: p for p in slide.shapes if p.shape_type == 13}
        self.assertEqual(set(pics), {"Picture TL", "Picture TR", "Picture B"})
        for name, box, ratio in (("Picture TL", BOX_TL, 1600 / 900),
                                 ("Picture TR", BOX_TR, 600 / 1200),
                                 ("Picture B", BOX_B, 1.0)):
            pic = pics[name]
            self.assertTrue(inside(pic, box), name)
            self.assertAlmostEqual(pic.width / pic.height, ratio, delta=ratio * 0.01)
            # contain: touches the box on at least one axis
            self.assertTrue(abs(pic.width - box[2]) <= 2 or abs(pic.height - box[3]) <= 2)

    def test_cover_fit_fills_box_and_blank_slot_is_dropped(self):
        slide = self.build().slides[2]
        pics = {p.name: p for p in slide.shapes if p.shape_type == 13}
        self.assertEqual(set(pics), {"Picture TL", "Picture B"})
        tl = pics["Picture TL"]
        self.assertEqual((tl.left, tl.top, tl.width, tl.height), BOX_TL)
        self.assertGreater(tl.crop_left + tl.crop_right + tl.crop_top + tl.crop_bottom, 0)

    def test_text_slot_keeps_formatting_and_static_shapes_copied(self):
        slide = self.build().slides[1]
        caption = next(s for s in slide.shapes if s.name == "Caption")
        self.assertEqual(caption.text_frame.text, "New caption")
        self.assertEqual(caption.text_frame.paragraphs[0].runs[0].font.size, Pt(14))
        self.assertTrue(any(s.name == "Rule" for s in slide.shapes))
        self.assertEqual(slide.notes_slide.notes_text_frame.text, "Speaker note")

    def test_unfilled_text_slot_is_dropped(self):
        slide = self.build().slides[2]
        self.assertFalse(any(s.name == "Caption" for s in slide.shapes))

    def test_shape_ids_unique(self):
        for slide in self.build().slides:
            ids = [s.shape_id for s in slide.shapes]
            self.assertEqual(len(ids), len(set(ids)))

    def test_too_many_images_is_an_error(self):
        with self.assertRaisesRegex(ValueError, "image"):
            self.build(SPEC.replace('"square.png"]\ntexts', '"square.png", "wide.png"]\ntexts'))

    def test_unknown_pattern_and_missing_file_are_errors(self):
        with self.assertRaisesRegex(ValueError, "pattern"):
            self.build(SPEC.replace('pattern = "three-up"\ntitle = "Loss', 'pattern = "nope"\ntitle = "Loss'))
        with self.assertRaisesRegex(FileNotFoundError, "tall.png"):
            (self.dir / "tall.png").unlink()
            self.build()

    def test_default_output_path_is_dated_in_task_out(self):
        expected = self.dir / "out" / f"template-{date.today():%y%m%d}.pptx"
        self.assertEqual(build_slides.default_output_path(self.spec, self.template), expected)

    def test_refuses_to_overwrite_template(self):
        with self.assertRaisesRegex(ValueError, "template"):
            build_slides.build(self.spec, self.template)
    def test_keep_true_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "out of range"):
            self.build(SPEC.replace("keep = 1", "keep = true"))

    def test_slide_hyperlink_to_dropped_slide_is_removed(self):
        prs = Presentation(self.template)
        caption = next(s for s in prs.slides[1].shapes if s.name == "Caption")
        caption.click_action.target_slide = prs.slides[0]  # link pattern -> cover
        prs.save(self.template)
        self.build(SPEC.replace("[[slides]]\nkeep = 1\n", ""))  # cover not kept
        import zipfile
        names = zipfile.ZipFile(self.dir / "out.pptx").namelist()
        self.assertEqual(len([n for n in names if re.match(r"ppt/slides/slide\d+\.xml$", n)]), 2)


class ReadingOrderTests(unittest.TestCase):
    def test_rows_then_left_to_right(self):
        class S:
            def __init__(self, n, x, y, w, h):
                self.name, self.left, self.top, self.width, self.height = n, x, y, w, h
        shapes = [S("TL", *BOX_TL), S("B", *BOX_B), S("TR", *BOX_TR)]
        self.assertEqual([s.name for s in slides_lib.reading_order(shapes)], ["TL", "TR", "B"])


if __name__ == "__main__":
    unittest.main()

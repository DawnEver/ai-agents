import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

import build_slides  # noqa: E402
import inspect_slides  # noqa: E402
import lint_slides  # noqa: E402
from fixtures import W, make_image, make_template  # noqa: E402


class InspectTests(unittest.TestCase):
    def test_reports_slots_in_reading_order(self):
        with tempfile.TemporaryDirectory() as d:
            info = inspect_slides.inspect(make_template(Path(d)))
        self.assertEqual(info["slide_count"], 2)
        slots = [(s["slot"], s["name"]) for s in info["slides"][1]["shapes"] if s["slot"]]
        # shapes are listed in z-order; the slot label carries the reading order
        self.assertEqual(dict(slots), {"title": "Title 1", "image[0]": "Picture TL",
                                       "image[1]": "Picture TR", "image[2]": "Picture B",
                                       "text[0]": "Caption"})
        md = inspect_slides.to_markdown(info)
        self.assertIn("image[1]", md)
        self.assertIn("## Slide 2", md)


class LintTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def codes(self, path):
        return {i.code for i in lint_slides.lint(path)}

    def test_flags_distortion_offslide_and_overlap(self):
        img = make_image(self.dir / "a.png", (1000, 500))
        prs = Presentation()
        prs.slide_width = Emu(W)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(img), 0, 0, Emu(1000000), Emu(1000000))       # 2:1 → 1:1 stretched
        slide.shapes.add_picture(str(img), Emu(W - 500000), 0, Emu(1000000), Emu(500000))  # off right edge
        slide.shapes.add_picture(str(img), Emu(500000), Emu(500000), Emu(1000000), Emu(500000))  # overlaps #1
        path = self.dir / "bad.pptx"
        prs.save(path)
        self.assertTrue({"distorted-image", "off-slide", "overlap"} <= self.codes(path))
        self.assertTrue(lint_slides.has_errors(lint_slides.lint(path)))

    def test_empty_placeholder_warning(self):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[1])  # title + body, both empty
        path = self.dir / "empty.pptx"
        prs.save(path)
        self.assertIn("empty-placeholder", self.codes(path))

    def test_clean_build_has_no_errors(self):
        template = make_template(self.dir)
        make_image(self.dir / "x.png", (800, 600))
        spec = self.dir / "deck.toml"
        spec.write_text('template = "template.pptx"\n[patterns.p]\nslide = 2\n'
                        '[[slides]]\npattern = "p"\ntitle = "T"\nimages = ["x.png"]\n')
        out = self.dir / "out.pptx"
        build_slides.build(spec, out)
        issues = lint_slides.lint(out)
        self.assertFalse(lint_slides.has_errors(issues), issues)
        self.assertTrue(template.exists())


class TableTests(unittest.TestCase):
    def test_empty_table_is_an_error(self):
        import slides_lib
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        frame = slide.shapes.add_table(2, 2, 0, 0, Emu(2000000), Emu(800000))
        with self.assertRaisesRegex(ValueError, "row"):
            slides_lib.set_table(frame._element, [])


if __name__ == "__main__":
    unittest.main()

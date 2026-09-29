import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tests"))

import md2slides  # noqa: E402
import slides2md  # noqa: E402
from fixtures import make_template  # noqa: E402


class TextRoundTripTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.src = make_template(self.dir)
        prs = Presentation(self.src)
        slide = prs.slides[0]
        box = slide.shapes.add_textbox(Emu(100000), Emu(2000000), Emu(5000000), Emu(2000000))
        box.name = "Bullets"
        tf = box.text_frame
        tf.text = "First point"
        tf.paragraphs[0].runs[0].font.size = Pt(20)
        p = tf.add_paragraph()
        p.text, p.level = "Sub point", 1
        table = slide.shapes.add_table(2, 2, Emu(6000000), Emu(2000000), Emu(4000000), Emu(800000)).table
        for r, row in enumerate([["Speed", "Eff"], ["1000", "93.2"]]):
            for c, v in enumerate(row):
                table.cell(r, c).text = v
        prs.save(self.src)

    def tearDown(self):
        self.tmp.cleanup()

    def test_extract_is_stable_and_complete(self):
        md = slides2md.to_markdown(self.src)
        self.assertIn("## Slide 1", md)
        self.assertIn("Cover", md)
        self.assertIn("First point\n  Sub point", md)
        self.assertIn("| Speed | Eff |", md)
        self.assertEqual(md, slides2md.to_markdown(self.src))

    def test_edit_apply_reextract(self):
        md = slides2md.to_markdown(self.src)
        edited = (md.replace("Cover", "New **cover**")
                    .replace("Sub point", "Changed sub")
                    .replace("| 1000 | 93.2 |", "| 1000 | 94.0 |\n| 2000 | 95.1 |"))
        md_path, out = self.dir / "deck.md", self.dir / "out.pptx"
        md_path.write_text(edited)
        md2slides.apply(md_path, self.src, out)
        again = slides2md.to_markdown(out)
        self.assertEqual(again.split("\n", 1)[1], edited.split("\n", 1)[1])  # line 1 names the source file

        prs = Presentation(out)
        box = next(s for s in prs.slides[0].shapes if s.name == "Bullets")
        self.assertEqual(box.text_frame.paragraphs[0].runs[0].font.size, Pt(20))
        self.assertEqual(box.text_frame.paragraphs[1].level, 1)

    def test_unknown_anchor_is_an_error(self):
        md_path = self.dir / "deck.md"
        md_path.write_text("<!-- s9999/1 -->\nText\n")
        with self.assertRaisesRegex(ValueError, "s9999/1"):
            md2slides.apply(md_path, self.src, self.dir / "out.pptx")


class RoundTripEdgeCaseTests(unittest.TestCase):
    def roundtrip(self, setup):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            setup(slide)
            src, md_path, out = d / "a.pptx", d / "a.md", d / "b.pptx"
            prs.save(src)
            md = slides2md.to_markdown(src)
            md_path.write_text(md)
            md2slides.apply(md_path, src, out)
            return md, slides2md.to_markdown(out), Presentation(out)

    def box(self, slide, text):
        tb = slide.shapes.add_textbox(Emu(0), Emu(0), Emu(3000000), Emu(2000000))
        tb.text_frame.text = text
        return tb

    def test_heading_like_lines_survive(self):
        md, again, prs = self.roundtrip(lambda s: self.box(s, "intro\n## Slide 9 recap\n<!-- s1/2 -->\nend"))
        self.assertEqual(prs.slides[0].shapes[0].text_frame.text, "intro\n## Slide 9 recap\n<!-- s1/2 -->\nend")
        self.assertEqual(md.split("\n", 1)[1], again.split("\n", 1)[1])

    def test_line_break_stays_a_line_break(self):
        md, again, prs = self.roundtrip(lambda s: self.box(s, "line1\vline2"))
        tf = prs.slides[0].shapes[0].text_frame
        self.assertEqual(len(tf.paragraphs), 1)
        self.assertEqual(tf.text, "line1\vline2")
        self.assertEqual(md.split("\n", 1)[1], again.split("\n", 1)[1])


if __name__ == "__main__":
    unittest.main()

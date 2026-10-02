"""Shared primitives: shape classification, reading order, slot discovery, text/table filling,
relationship-aware shape copying. Pure python-pptx + lxml; no PowerPoint needed."""
import copy
import re
import sys

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

if sys.version_info >= (3, 11):
    import tomllib
else:  # pragma: no cover
    import tomli as tomllib

EMU_PER_INCH = 914400
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
FOOTER_TYPES = {PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE}
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
BR = "<br>"  # in text: a line break (a:br) inside a paragraph


def load_toml(path):
    with open(path, "rb") as f:
        return tomllib.load(f)


def inches(emu):
    return round((emu or 0) / EMU_PER_INCH, 2)


# --- classification -------------------------------------------------------

def placeholder_type(shape):
    return shape.placeholder_format.type if shape.is_placeholder else None


def classify(shape, static_names=()):
    """Return one of: title, footer, picture, table, text, static."""
    if shape.name in static_names:
        return "static"
    ph = placeholder_type(shape)
    if ph in TITLE_TYPES:
        return "title"
    if ph in FOOTER_TYPES:
        return "footer"
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or ph == PP_PLACEHOLDER.PICTURE:
        return "picture"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if shape.shape_type != MSO_SHAPE_TYPE.GROUP and shape.has_text_frame and shape.text_frame.text.strip():
        return "text"
    return "static"


SLOT_PREFIX = {"picture": "image", "text": "text", "table": "table"}


def reading_order(shapes):
    """Rows top→bottom (a shape joins a row when its top is within half a height of the row's
    first shape), then left→right inside each row."""
    rows = []
    for s in sorted(shapes, key=lambda s: (s.top or 0, s.left or 0)):
        if rows:
            head = rows[-1][0]
            if (s.top or 0) - (head.top or 0) < 0.5 * min(head.height or 1, s.height or 1):
                rows[-1].append(s)
                continue
        rows.append([s])
    return [s for row in rows for s in sorted(row, key=lambda s: s.left or 0)]


def slot_map(slide, static_names=()):
    """{shape_id: slot label} for fillable shapes, e.g. {5: 'image[0]', 2: 'title'}."""
    labels, by_kind = {}, {}
    for shape in slide.shapes:
        kind = classify(shape, static_names)
        if kind == "title":
            labels[shape.shape_id] = "title"
        elif kind in SLOT_PREFIX:
            by_kind.setdefault(kind, []).append(shape)
    for kind, shapes in by_kind.items():
        for i, s in enumerate(reading_order(shapes)):
            labels[s.shape_id] = f"{SLOT_PREFIX[kind]}[{i}]"
    return labels


# --- copying --------------------------------------------------------------

def append_element(slide, el):
    slide.shapes._spTree.insert_element_before(el, "p:extLst")
    return el


def copy_shape(el, src_part, dst_slide):
    """Deep-copy a shape element onto dst_slide, re-pointing every relationship it uses
    (any r:* attribute — pictures, links, charts, SmartArt). Links to other slides are removed:
    the target may be dropped from the output, and a kept link would drag it back in."""
    new = copy.deepcopy(el)
    dst_part = dst_slide.part
    for node in list(new.iter()):
        for attr, rid in list(node.attrib.items()):
            if not attr.startswith(R_NS) or rid not in src_part.rels:
                continue
            rel = src_part.rels[rid]
            if rel.reltype == RT.SLIDE:
                if node.getparent() is not None:
                    node.getparent().remove(node)
                break
            if rel.is_external:
                new_rid = dst_part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                new_rid = dst_part.relate_to(rel.target_part, rel.reltype)
            node.set(attr, new_rid)
    return append_element(dst_slide, new)


def renumber_shape_ids(slide):
    """Make shape ids unique; connector end-points follow their shapes."""
    tree = slide.shapes._spTree
    mapping = {}
    for i, node in enumerate(tree.iter(qn("p:cNvPr")), start=2):
        mapping.setdefault(node.get("id"), str(i))
        node.set("id", str(i))
    for cxn in tree.iter(qn("a:stCxn"), qn("a:endCxn")):
        if cxn.get("id") in mapping:
            cxn.set("id", mapping[cxn.get("id")])


# --- text -----------------------------------------------------------------

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def split_bold(line):
    parts, pos = [], 0
    for m in BOLD_RE.finditer(line):
        if m.start() > pos:
            parts.append((line[pos:m.start()], False))
        parts.append((m.group(1), True))
        pos = m.end()
    if pos < len(line) or not parts:
        parts.append((line[pos:], False))
    return parts


def _as_rpr(el):
    if el is None:
        return etree.Element(qn("a:rPr"))
    rpr = etree.Element(qn("a:rPr"), attrib=dict(el.attrib))
    for child in el:
        rpr.append(copy.deepcopy(child))
    return rpr


def set_text(tx_body, text):
    """Replace all paragraphs of an a:txBody / p:txBody with `text`, keeping the first paragraph's
    paragraph and run formatting. Two leading spaces per indent level; **bold** supported."""
    paras = tx_body.findall(qn("a:p"))
    proto = paras[0] if paras else etree.Element(qn("a:p"))
    ppr = proto.find(qn("a:pPr"))
    first_run = proto.find(qn("a:r"))
    rpr = _as_rpr(first_run.find(qn("a:rPr")) if first_run is not None else proto.find(qn("a:endParaRPr")))
    rpr.attrib.pop("b", None)
    for p in paras:
        tx_body.remove(p)
    for line in text.split("\n"):
        stripped = line.lstrip(" ")
        level = (len(line) - len(stripped)) // 2
        p = etree.SubElement(tx_body, qn("a:p"))
        if ppr is not None or level:
            new_ppr = copy.deepcopy(ppr) if ppr is not None else etree.Element(qn("a:pPr"))
            if level:
                new_ppr.set("lvl", str(level))
            else:
                new_ppr.attrib.pop("lvl", None)
            p.append(new_ppr)
        if not stripped:
            end = copy.deepcopy(rpr)
            end.tag = qn("a:endParaRPr")
            p.append(end)
            continue
        for seg, bold in split_bold(stripped):
            for k, piece in enumerate(seg.split(BR)):
                if k:
                    etree.SubElement(p, qn("a:br")).append(copy.deepcopy(rpr))
                if not piece:
                    continue
                r = etree.SubElement(p, qn("a:r"))
                run_rpr = copy.deepcopy(rpr)
                if bold:
                    run_rpr.set("b", "1")
                r.append(run_rpr)
                etree.SubElement(r, qn("a:t")).text = piece


def get_text(tx_body):
    """Inverse of set_text (indent levels and fully-bold runs)."""
    lines = []
    for p in tx_body.findall(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        level = int(ppr.get("lvl", "0")) if ppr is not None else 0
        chunks = []
        for r in p.iter(qn("a:r"), qn("a:fld"), qn("a:br")):
            if r.tag == qn("a:br"):
                chunks.append((BR, None))
                continue
            t = r.findtext(qn("a:t")) or ""
            rp = r.find(qn("a:rPr"))
            chunks.append((t, rp is not None and rp.get("b") in ("1", "true")))
        merged = []
        for t, b in chunks:  # merge adjacent runs with equal boldness
            if merged and b is None:  # a break joins whichever run precedes it
                merged[-1] = (merged[-1][0] + t, merged[-1][1])
            elif merged and merged[-1][1] == b:
                merged[-1] = (merged[-1][0] + t, b)
            else:
                merged.append((t, b))
        body = "".join(f"**{t}**" if b and t.strip() else t for t, b in merged)
        lines.append("  " * level + body if body else "")
    return "\n".join(lines)


def shape_tx_body(el):
    return el.find(qn("p:txBody"))


# --- tables ---------------------------------------------------------------

def table_rows(graphic_frame_el):
    return graphic_frame_el.findall(".//" + qn("a:tr"))


def get_table(graphic_frame_el):
    return [[get_text(tc.find(qn("a:txBody"))) for tc in tr.findall(qn("a:tc"))]
            for tr in table_rows(graphic_frame_el)]


def set_table(graphic_frame_el, rows):
    """Fill a table; row count follows `rows` (extra rows deep-copy the last row), columns are fixed."""
    trs = table_rows(graphic_frame_el)
    if not rows:
        raise ValueError("table needs at least one row (template content would otherwise remain)")
    if not trs:
        raise ValueError("table has no rows")
    ncols = len(trs[0].findall(qn("a:tc")))
    while len(trs) < len(rows):
        new = copy.deepcopy(trs[-1])
        trs[-1].addnext(new)
        trs.append(new)
    while len(trs) > len(rows) and len(trs) > 1:
        tr = trs.pop()
        tr.getparent().remove(tr)
    for tr, values in zip(trs, rows):
        if len(values) > ncols:
            raise ValueError(f"table row has {len(values)} cells but the table has {ncols} columns")
        for tc, value in zip(tr.findall(qn("a:tc")), list(values) + [""] * (ncols - len(values))):
            set_text(tc.find(qn("a:txBody")), str(value))

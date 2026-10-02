#!/usr/bin/env python3
"""Deterministic layout checks for a deck.

Usage:
    python scripts/lint_slides.py <deck.pptx> [--strict]

Errors (exit 1): distorted-image (aspect differs >1% from the source pixels), off-slide.
Warnings (exit 1 with --strict): overlap, empty-placeholder, text-overflow (estimate).
Lint is a floor, not a verdict — always look at the preview as well.
"""
import argparse
import io
import math
import sys
import unicodedata
from dataclasses import dataclass

from PIL import Image
from pptx import Presentation
from pptx.util import Pt

import slides_lib as lib

ASPECT_TOL = 0.01
EDGE_TOL = 0.005
OVERLAP_TOL = 0.02
ERRORS = {"distorted-image", "off-slide"}


@dataclass
class Issue:
    slide: int
    code: str
    shape: str
    detail: str

    @property
    def severity(self):
        return "error" if self.code in ERRORS else "warning"

    def __str__(self):
        return f"slide {self.slide}: {self.severity} {self.code} [{self.shape}] {self.detail}"


def has_errors(issues):
    return any(i.severity == "error" for i in issues)


def _box(s):
    return s.left or 0, s.top or 0, s.width or 0, s.height or 0


def _distortion(pic):
    try:
        if pic.image is None:
            return None
        with Image.open(io.BytesIO(pic.image.blob)) as im:
            iw, ih = im.size
    except Exception:  # vector (EMF/SVG) or unreadable: nothing to measure
        return None
    vis_w = iw * (1 - pic.crop_left - pic.crop_right)
    vis_h = ih * (1 - pic.crop_top - pic.crop_bottom)
    if not (vis_w > 0 and vis_h > 0 and pic.width and pic.height):
        return None
    return abs((pic.width / pic.height) / (vis_w / vis_h) - 1)


def _text_overflow(shape):
    tf = shape.text_frame
    body_pr = tf._txBody.find(lib.qn("a:bodyPr"))
    if body_pr is not None and (body_pr.find(lib.qn("a:normAutofit")) is not None
                                or body_pr.find(lib.qn("a:spAutoFit")) is not None):
        return None  # PowerPoint shrinks the text / grows the box
    width = (shape.width or 0) - Pt(14)  # default insets 0.1in each side
    if width <= 0:
        return None
    need = 0
    for p in tf.paragraphs:
        sizes = [r.font.size for r in p.runs if r.font.size] or [Pt(18)]
        size = max(sizes)
        ems = sum(1.0 if unicodedata.east_asian_width(ch) in "WF" else 0.5 for ch in p.text) or 0.5
        lines = max(1, math.ceil(ems * size / width))
        need += lines * size * 1.2
    have = (shape.height or 0) - Pt(7)
    return need / have if have > 0 and need > have * 1.05 else None


def lint(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    issues = []
    for n, slide in enumerate(prs.slides, start=1):
        content = []
        for s in slide.shapes:
            kind = lib.classify(s)
            x, y, w, h = _box(s)
            if w and h and (x < -sw * EDGE_TOL or y < -sh * EDGE_TOL
                            or x + w > sw * (1 + EDGE_TOL) or y + h > sh * (1 + EDGE_TOL)):
                issues.append(Issue(n, "off-slide", s.name, "extends beyond the slide"))
            if kind == "picture" and hasattr(s, "image"):
                d = _distortion(s)
                if d is not None and d > ASPECT_TOL:
                    issues.append(Issue(n, "distorted-image", s.name, f"aspect off by {d:.1%}"))
            if s.is_placeholder and kind != "footer" and s.has_text_frame and not s.text_frame.text.strip():
                issues.append(Issue(n, "empty-placeholder", s.name, "shows prompt text in edit view"))
            if kind == "text" or (kind == "title" and s.text_frame.text.strip()):
                r = _text_overflow(s)
                if r:
                    issues.append(Issue(n, "text-overflow", s.name, f"~{r:.0%} of box height needed"))
            if kind in ("picture", "text", "table", "title") and w and h:
                content.append(s)
        for i, a in enumerate(content):
            ax, ay, aw, ah = _box(a)
            for b in content[i + 1:]:
                bx, by, bw, bh = _box(b)
                ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
                iy = max(0, min(ay + ah, by + bh) - max(ay, by))
                frac = ix * iy / min(aw * ah, bw * bh)
                if frac > OVERLAP_TOL:
                    issues.append(Issue(n, "overlap", f"{a.name} × {b.name}", f"{frac:.0%} of the smaller"))
    return issues


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--strict", action="store_true", help="warnings also fail")
    args = ap.parse_args(argv)
    issues = lint(args.deck)
    for i in issues:
        print(i)
    errors = sum(i.severity == "error" for i in issues)
    print(f"{errors} error(s), {len(issues) - errors} warning(s)")
    sys.exit(1 if errors or (args.strict and issues) else 0)


if __name__ == "__main__":
    main()

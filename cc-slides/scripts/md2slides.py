#!/usr/bin/env python3
"""Patch edited Markdown (from slides2md.py) back into the deck it came from.

Usage:
    python scripts/md2slides.py <deck.md> <source.pptx> [output.pptx]

Only anchored blocks are applied; each replaces the text of that shape (first-run formatting,
paragraph formatting and geometry preserved). Table blocks may change row count; columns are
fixed. Unknown anchors are an error. Text outside any anchor is ignored with a warning — layout
changes belong in deck.toml, not here. Writes a new file; never overwrites the source.
"""
import argparse
import re
import sys
from datetime import date
from pathlib import Path

from pptx import Presentation

import slides_lib as lib

ANCHOR_RE = re.compile(r"^<!-- s(\d+)/(\d+|notes) -->$")
HEADING_RE = re.compile(r"^## Slide \d+")


def parse(md):
    """[(slide_id, shape_id|'notes', text)] plus a list of stray non-empty lines."""
    blocks, stray, current = [], [], None
    for line in md.split("\n"):
        m = ANCHOR_RE.match(line)
        if m:
            current = [int(m.group(1)), m.group(2), []]
            blocks.append(current)
        elif HEADING_RE.match(line) or line.startswith("<!-- cc-slides source:"):
            current = None
        elif current is not None:
            current[2].append(line[1:] if line.startswith("\\") else line)
        elif line.strip():
            stray.append(line)
    parsed = []
    for sid, shid, lines in blocks:
        while lines and not lines[-1].strip():
            lines.pop()
        parsed.append((sid, shid if shid == "notes" else int(shid), "\n".join(lines)))
    return parsed, stray


def _split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip())[1:-1]
    return [c.strip().replace("\\|", "|").replace("<p>", "\n") for c in cells]


def parse_table(text):
    rows = [_split_row(l) for l in text.split("\n") if l.strip().startswith("|")]
    return [r for i, r in enumerate(rows) if not (i == 1 and all(set(c) <= set("-: ") for c in r))]


def default_output_path(md_path, source):
    return Path(md_path).parent / "out" / f"{Path(source).stem}-{date.today():%y%m%d}.pptx"


def apply(md_path, source, output=None):
    source = Path(source).resolve()
    out = Path(output) if output else default_output_path(md_path, source)
    if out.resolve() == source:
        raise ValueError("refusing to overwrite the source deck; choose another output path")
    blocks, stray = parse(Path(md_path).read_text(encoding="utf-8"))
    for line in stray:
        print(f"warning: ignored text outside any anchor: {line[:60]!r}", file=sys.stderr)

    prs = Presentation(source)
    for sid, shid, text in blocks:
        slide = prs.slides.get(sid)
        if slide is None:
            raise ValueError(f"anchor s{sid}/{shid}: no such slide")
        if shid == "notes":
            slide.notes_slide.notes_text_frame.text = text
            continue
        shape = next((s for s in slide.shapes if s.shape_id == shid), None)
        if shape is None:
            raise ValueError(f"anchor s{sid}/{shid}: no such shape")
        if getattr(shape, "has_table", False) and shape.has_table:
            lib.set_table(shape._element, parse_table(text))
        elif shape.has_text_frame:
            lib.set_text(lib.shape_tx_body(shape._element), text)
        else:
            raise ValueError(f"anchor s{sid}/{shid}: shape {shape.name!r} holds no text")
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("markdown")
    ap.add_argument("source")
    ap.add_argument("output", nargs="?")
    args = ap.parse_args(argv)
    try:
        print(apply(args.markdown, args.source, args.output))
    except ValueError as e:
        sys.exit(f"error: {e}")


if __name__ == "__main__":
    main()

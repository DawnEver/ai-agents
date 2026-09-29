#!/usr/bin/env python3
"""Transcribe every text-bearing shape of a deck to anchored Markdown, for wording edits.

Usage:
    python scripts/slides2md.py <deck.pptx> [output.md]

Each block starts with `<!-- s<slide_id>/<shape_id> -->` (the round-trip map; md2slides.py
patches text back by these ids). Text keeps indent levels (two spaces per level) and **bold**
runs; tables become Markdown tables. Geometry is not part of this format — use deck.toml for layout.
Never modifies the input deck.
"""
import argparse
import re
import sys
from pathlib import Path

from pptx import Presentation

import slides_lib as lib


ESCAPE_RE = re.compile(r"^(## Slide|<!--|\\)")


def _escape(cell):
    return cell.replace("|", "\\|").replace("\n", "<p>")


def _escape_lines(text):
    """Lines that look like md structure get a leading backslash (md2slides strips it)."""
    return "\n".join("\\" + l if ESCAPE_RE.match(l) else l for l in text.split("\n"))


def _table_md(rows):
    lines = ["| " + " | ".join(_escape(c) for c in rows[0]) + " |",
             "|" + "---|" * len(rows[0])]
    lines += ["| " + " | ".join(_escape(c) for c in row) + " |" for row in rows[1:]]
    return "\n".join(lines)


def _blocks(slide):
    for shape in lib.reading_order(list(slide.shapes)):
        kind = lib.classify(shape)
        anchor = f"<!-- s{slide.slide_id}/{shape.shape_id} -->"
        if kind == "table":
            yield anchor, _table_md(lib.get_table(shape._element))
        elif kind in ("title", "text") and shape.has_text_frame:
            yield anchor, _escape_lines(lib.get_text(lib.shape_tx_body(shape._element)))


def to_markdown(path):
    prs = Presentation(path)
    out = [f"<!-- cc-slides source: {Path(path).name} -->"]
    for n, slide in enumerate(prs.slides, start=1):
        out += ["", f"## Slide {n} · {slide.slide_layout.name}"]
        for anchor, body in _blocks(slide):
            out += ["", anchor, body]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out += ["", f"<!-- s{slide.slide_id}/notes -->",
                    _escape_lines(slide.notes_slide.notes_text_frame.text)]
    return "\n".join(out) + "\n"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("output", nargs="?")
    args = ap.parse_args(argv)
    md = to_markdown(args.deck)
    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(md, encoding="utf-8")
        print(args.output)
    else:
        sys.stdout.write(md)


if __name__ == "__main__":
    main()

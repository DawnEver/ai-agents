#!/usr/bin/env python3
"""Describe a deck: slide size, layouts, and every shape with its geometry and slot label.

Usage:
    python scripts/inspect_slides.py <deck.pptx> [--slides 3,19] [--json]

Slot labels (title, image[0], text[1], table[0] …) are exactly the order build_slides.py fills a
pattern in: reading order — rows top→bottom, then left→right. `static` shapes are copied verbatim.
"""
import argparse
import json
import sys

from pptx import Presentation

import slides_lib as lib


def inspect(path, only=None):
    prs = Presentation(path)
    info = {
        "path": str(path),
        "slide_width_in": lib.inches(prs.slide_width),
        "slide_height_in": lib.inches(prs.slide_height),
        "slide_count": len(prs.slides),
        "layouts": [layout.name for layout in prs.slide_layouts],
        "slides": [],
    }
    for n, slide in enumerate(prs.slides, start=1):
        if only and n not in only:
            continue
        labels = lib.slot_map(slide)
        shapes = []
        for shape in slide.shapes:
            text = shape.text_frame.text if shape.has_text_frame else ""
            shapes.append({
                "id": shape.shape_id,
                "name": shape.name,
                "kind": lib.classify(shape),
                "slot": labels.get(shape.shape_id),
                "left_in": lib.inches(shape.left), "top_in": lib.inches(shape.top),
                "width_in": lib.inches(shape.width), "height_in": lib.inches(shape.height),
                "text": " / ".join(text.split("\n"))[:80],
            })
        info["slides"].append({"index": n, "slide_id": slide.slide_id,
                               "layout": slide.slide_layout.name, "shapes": shapes})
    return info


def to_markdown(info):
    out = [f"# {info['path']}", "",
           f"{info['slide_width_in']} x {info['slide_height_in']} in, {info['slide_count']} slides", "",
           "## Layouts", ""]
    out += [f"{i}. {name}" for i, name in enumerate(info["layouts"], start=1)]
    for s in info["slides"]:
        out += ["", f"## Slide {s['index']} · {s['layout']}", "",
                "| slot | name | kind | left | top | width | height | text |",
                "|---|---|---|---|---|---|---|---|"]
        for sh in s["shapes"]:
            text = sh["text"].replace("|", "\\|")
            out.append(f"| {sh['slot'] or ''} | {sh['name']} | {sh['kind']} | {sh['left_in']} | "
                       f"{sh['top_in']} | {sh['width_in']} | {sh['height_in']} | {text} |")
    return "\n".join(out) + "\n"


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("deck")
    ap.add_argument("--slides", help="comma-separated 1-based slide numbers")
    ap.add_argument("--json", action="store_true")
    args = ap.parse_args(argv)
    only = {int(x) for x in args.slides.split(",")} if args.slides else None
    info = inspect(args.deck, only)
    sys.stdout.write(json.dumps(info, indent=2, ensure_ascii=False) + "\n" if args.json else to_markdown(info))


if __name__ == "__main__":
    main()
